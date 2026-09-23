"""Chargeur des fichiers clusters du Bilan (modèle « Template pour collecte de données mensuel »).
Lit, pour un dossier mensuel <mois>/Clusters :
  - les feuilles « Indicateur clé » -> hpc_indicator_values (réalisé cumulé par province)
  - la feuille « Case load Monitoring template » -> cluster_reached_disagg (statut, sexe/âge par province)
Le cluster vient du nom de fichier ; l'indicateur vient du libellé en A1 (les noms d'onglet ne sont pas fiables).
Usage : python charger_clusters.py <dossier> <mois> <sortie.sql>"""
import re, sys, glob, os, unicodedata
from openpyxl import load_workbook

def nrm(s): return unicodedata.normalize('NFKD', str(s or '')).encode('ascii','ignore').decode().lower()

FILE_CLUSTER=[('secal','SECAL'),('wash','WASH'),('sant','SANTE'),('protection','PRO'),('educ','EDU'),('edu-','EDU'),
              ('nut','NUT'),('abris','ABRIS'),('gsat','GSAT'),('multisecteur','refugee'),('refugi','refugee')]
IND_RULES={
 'SECAL':[('alimentaire','SECAL-01'),('animaux','SECAL-02'),('maraich','SECAL-03')],
 'WASH': [('latrine','WASH-02'),('15 l','WASH-01'),('eau aux normes','WASH-01'),('promotion','WASH-03')],
 'SANTE':[('consultation','SANTE-01'),('accouchement','SANTE-02'),('pcime','SANTE-03'),('agents communautaire','SANTE-03')],
 'PRO':  [('ltb','PRO-04'),('assistance individuelle','PRO-02'),('bien etre','PRO-03'),('psychosocial','PRO-01')],
 'EDU':  [('fourniture','EDU-02'),('psychosocial','EDU-03'),('education formelle','EDU-01'),('formelle','EDU-01')],
 'NUT':  [('enceintes','NUT-03'),('moderee','NUT-02'),('(mam)','NUT-02'),('severe','NUT-01'),('(mas)','NUT-01')],
 'ABRIS':[('kits','ABRIS-01'),('urgence','ABRIS-02'),('transitionnel','ABRIS-03')],
 'GSAT': [('profilage','GSAT-01'),('coordination','GSAT-02'),('plainte','GSAT-03')],
 'refugee':[('attestation','REF-01')],
}
DIS_COLS={'non-deplaces':'non_deplaces','pdi':'pdi','retournes':'retournes','refugies':'refugies','filles':'filles','garcons':'garcons',
          'femmes':'femmes','hommes':'hommes','personne agees femmes':'agees_f','personne agees hommes':'agees_h',
          'personnes handicapees':'handicapes','atteint cumulatif':'total'}

def cluster_of(fn):
    f=nrm(os.path.basename(fn))
    for k,c in FILE_CLUSTER:
        if k in f: return c
    return None

def header(ws):
    for r in range(1,8):
        vals=[nrm(c.value).strip() for c in ws[r]]
        for i,v in enumerate(vals):
            if 'pcode' in v and ('amd2' in v or 'adm2' in v): return r,vals,i
    return None,None,None

def provinces(ws,hr,pc):
    for row in ws.iter_rows(min_row=hr+1,values_only=True):
        p=row[pc] if pc<len(row) else None
        if isinstance(p,str) and re.match(r'^BF\d{4}$',p.strip()): yield p.strip(),row

def num(v): return round(v) if isinstance(v,(int,float)) else None

def read_file(fn):
    cl=cluster_of(fn); wb=load_workbook(fn,data_only=True); inds={}; cibles={}; dis=[]; warn=[]
    # plusieurs feuilles « Case load » possibles (ex. Santé : feuille du mois seul) : garder la feuille cumulative « Case load Monitoring template »
    cls=[ws for ws in wb.worksheets if ws.sheet_state=='visible' and 'case load' in ws.title.lower()]
    keep=next((ws for ws in cls if nrm(ws.title).strip()=='case load monitoring template'),cls[0] if cls else None)
    for ws in wb.worksheets:
        if 'case load' in ws.title.lower() and ws is not keep: warn.append(f'feuille ignorée : {ws.title}'); continue
        if ws.sheet_state!='visible': continue
        hr,vals,pc=header(ws)
        if hr is None: continue
        if 'case load' in ws.title.lower():
            cols={DIS_COLS[v]:i for i,v in enumerate(vals) if v in DIS_COLS}
            for p,row in provinces(ws,hr,pc):
                dis.append((p,{k:num(row[i]) for k,i in cols.items()}))
            continue
        a1=nrm(ws['A1'].value); code=next((c for k,c in IND_RULES.get(cl,[]) if k in a1),None)
        if not code: warn.append(f'onglet non reconnu : {ws.title} | {str(ws["A1"].value)[:80]}'); continue
        if code in inds: warn.append(f'{code} en double ({ws.title}) — premier gardé'); continue
        cu=next((i for i,v in enumerate(vals) if 'cumul' in v),pc+1)
        ci=next((i for i,v in enumerate(vals) if v=='cible'),None)
        inds[code]={p:num(row[cu]) for p,row in provinces(ws,hr,pc)}
        if ci is not None: cibles[code]={p:num(row[ci]) for p,row in provinces(ws,hr,pc)}
    return cl,inds,dis,warn,cibles

def pick_files(folder):
    """Un fichier par cluster ; si une version « revu » existe, elle remplace l'autre."""
    by={}
    for f in sorted(glob.glob(os.path.join(folder,'*.xlsx'))):
        if os.path.basename(f).startswith('~$') or '3w' in nrm(os.path.basename(f)): continue
        c=cluster_of(f)
        if not c: continue
        if c not in by or 'revu' in nrm(f): by[c]=f
    return by

def q(s): return "'"+str(s).replace("'","''")+"'"
def nv(v): return 'null' if v is None else str(v)

def build_sql(folder,mois):
    ind_rows=[]; dis_rows=[]; report=[]
    for cl,f in sorted(pick_files(folder).items()):
        c,inds,dis,warn,_=read_file(f); src=os.path.basename(f)
        for code,vals in sorted(inds.items()):
            for p,v in vals.items(): ind_rows.append(f"({q(code)},{mois},{q(p)},{nv(v)},{q(src)})")
            report.append((cl,code,sum(v or 0 for v in vals.values())))
        for p,d in dis:
            dis_rows.append(f"({q(cl)},{mois},{q(p)},"+','.join(nv(d.get(k)) for k in ['non_deplaces','pdi','retournes','refugies','filles','garcons','femmes','hommes','agees_f','agees_h','handicapes','total'])+f",{q(src)})")
        tot=sum((d.get('total') or 0) for _,d in dis); report.append((cl,'CASELOAD',tot))
        for w in warn: report.append((cl,'ATTENTION',w))
    sql=[]
    if ind_rows: sql.append("insert into hpc_indicator_values (indicator_code,mois,adm2_pcode,realise,source_file) values\n"+',\n'.join(ind_rows)+
        "\non conflict (hpc_cycle,indicator_code,mois,adm2_pcode) do update set realise=excluded.realise, source_file=excluded.source_file, loaded_at=now();")
    if dis_rows: sql.append("insert into cluster_reached_disagg (cluster,mois,adm2_pcode,non_deplaces,pdi,retournes,refugies,filles,garcons,femmes,hommes,agees_f,agees_h,handicapes,total,source_file) values\n"+',\n'.join(dis_rows)+
        "\non conflict (hpc_cycle,cluster,mois,adm2_pcode) do update set non_deplaces=excluded.non_deplaces,pdi=excluded.pdi,retournes=excluded.retournes,refugies=excluded.refugies,filles=excluded.filles,garcons=excluded.garcons,femmes=excluded.femmes,hommes=excluded.hommes,agees_f=excluded.agees_f,agees_h=excluded.agees_h,handicapes=excluded.handicapes,total=excluded.total,source_file=excluded.source_file,loaded_at=now();")
    return sql,report

if __name__=='__main__':
    folder,mois,out=sys.argv[1],int(sys.argv[2]),sys.argv[3]
    sql,report=build_sql(folder,mois)
    for r in report: print(*r,sep=' | ')
    for i,s in enumerate(sql): open(f'{out}.{i}.sql','w').write(s)
