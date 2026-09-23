"""Lit un modèle indicateurs rempli -> lignes (indicator_code, mois, adm2_pcode, cible, realise) + SQL upsert."""
import sys, re
from openpyxl import load_workbook
def read(path):
    wb=load_workbook(path,data_only=True); mois=wb['Lisez-moi']['B3'].value
    assert isinstance(mois,(int,float)) and 1<=mois<=12, 'Mois manquant (Lisez-moi!B3)'
    rows=[]
    for ws in wb.worksheets:
        if ws.title in ('Lisez-moi','Catalogue'): continue
        codes={c.column:c.value for c in ws[1] if isinstance(c.value,str) and re.match(r'^[A-Z]+-\d\d$',c.value)}
        for r in range(4,ws.max_row+1):
            pc=ws.cell(row=r,column=3).value
            if not (isinstance(pc,str) and pc.startswith('BF')): continue
            for col,code in codes.items():
                cb,re_=ws.cell(row=r,column=col).value,ws.cell(row=r,column=col+1).value
                if cb is None and re_ is None: continue
                rows.append((code,int(mois),pc,None if cb is None else int(cb),None if re_ is None else int(re_)))
    return rows
def sql(rows,src):
    v=',\n'.join(f"('{c}',{m},'{p}',{'null' if cb is None else cb},{'null' if r is None else r},'{src}')" for c,m,p,cb,r in rows)
    return f"""insert into hpc_indicator_values (indicator_code,mois,adm2_pcode,cible,realise,source_file) values
{v}
on conflict (hpc_cycle,indicator_code,mois,adm2_pcode) do update set cible=excluded.cible, realise=excluded.realise, source_file=excluded.source_file, loaded_at=now();"""
if __name__=='__main__':
    rows=read(sys.argv[1]); print(len(rows),'lignes'); open(sys.argv[2],'w').write(sql(rows,sys.argv[1].split('/')[-1]))
